from django.shortcuts import render, redirect
from django.contrib import messages
import base64
import binascii
from django.utils import timezone
from io import BytesIO
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db import transaction
from django.contrib.auth import (authenticate,login,logout)
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.contrib.auth.password_validation import validate_password
from django.db import transaction
from django.conf import settings
from django.core.exceptions import ValidationError
from django.db.models import Avg, Q
from django.core.paginator import Paginator
from django.http import JsonResponse, HttpResponse
from .forms import (ReviewForm,FacultyProfileForm,EducationFormSet,SkillFormSet)
from .models import (Review,FacultyProfile)
from .email_utils import (send_student_thank_you,send_admin_notification)
from .sentiment_utils import analyze_sentiment
from .google_auth import create_google_flow
from google.oauth2 import id_token
from google.auth.transport import requests as google_requests
import requests
import csv
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import (Font,PatternFill,Alignment)
# =========================================================
# HOME
# =========================================================

def home(request):
    profile = (
        FacultyProfile.objects
        .prefetch_related("educations", "skills")
        .first()
    )

    return render(
        request,
        "home.html",
        {
            "profile": profile,
        }
    )

# =========================================================
# SUBMIT REVIEW
# =========================================================

def submit_review(request):

    if request.method == 'POST':

        form = ReviewForm(request.POST)

        if form.is_valid():

            # Create review without saving first
            review = form.save(commit=False)


            # Analyze sentiment
            review.sentiment = analyze_sentiment(
                review.feedback
            )


            # Save review
            review.save()


            print(
                f"Sentiment detected: {review.sentiment}"
            )


            # =================================================
            # STUDENT THANK YOU EMAIL
            # =================================================

            try:

                send_student_thank_you(review)

                print(
                    "Student email sent successfully."
                )

            except Exception as e:

                print(
                    "Student email error:",
                    e
                )


            # =================================================
            # ADMIN NOTIFICATION EMAIL
            # =================================================

            try:

                send_admin_notification(review)

                print(
                    "Admin notification sent successfully."
                )

            except Exception as e:

                print(
                    "Admin email error:",
                    e
                )


            return redirect(
                'thank_you'
            )


        else:

            messages.error(
                request,
                'Please correct the errors in the form.'
            )


    else:

        form = ReviewForm()


    return render(
        request,
        'review_form.html',
        {
            'form': form
        }
    )


# =========================================================
# THANK YOU
# =========================================================

def thank_you(request):

    return render(
        request,
        'thank_you.html'
    )


# =========================================================
# ADMIN LOGIN
# =========================================================

def admin_login(request):

    # Already logged in
    if request.user.is_authenticated:

        if request.user.is_staff:

            return redirect(
                'admin_dashboard'
            )

        logout(request)


    if request.method == 'POST':

        username = request.POST.get(
            'username'
        )

        password = request.POST.get(
            'password'
        )


        user = authenticate(
            request,
            username=username,
            password=password
        )


        if user is not None and user.is_staff:

            login(
                request,
                user
            )

            return redirect(
                'admin_dashboard'
            )


        messages.error(
            request,
            'Invalid admin credentials.'
        )


    return render(
        request,
        'admin_login.html'
    )


# =========================================================
# ADMIN LOGOUT
# =========================================================

@login_required(login_url='admin_login')
def admin_logout(request):

    logout(request)

    return redirect(
        'admin_login'
    )


# =========================================================
# FILTER REVIEWS
# =========================================================

def get_filtered_reviews(request):

    reviews = Review.objects.all().order_by(
        '-created_at'
    )


    # =====================================================
    # SEARCH
    # =====================================================

    search_query = request.GET.get(
        'search',
        ''
    ).strip()


    if search_query:

        reviews = reviews.filter(

            Q(
                student_name__icontains=
                search_query
            )

            |

            Q(
                student_email__icontains=
                search_query
            )

            |

            Q(
                course__icontains=
                search_query
            )

            |

            Q(
                batch__icontains=
                search_query
            )

            |

            Q(
                feedback__icontains=
                search_query
            )

        )


    # =====================================================
    # COURSE
    # =====================================================

    selected_course = request.GET.get(
        'course',
        ''
    ).strip()


    if selected_course:

        reviews = reviews.filter(
            course=selected_course
        )


    # =====================================================
    # RATING
    # =====================================================

    selected_rating = request.GET.get(
        'rating',
        ''
    ).strip()


    if selected_rating:

        reviews = reviews.filter(
            rating=selected_rating
        )


    # =====================================================
    # SENTIMENT
    # =====================================================

    selected_sentiment = request.GET.get(
        'sentiment',
        ''
    ).strip()


    if selected_sentiment:

        reviews = reviews.filter(
            sentiment=selected_sentiment
        )


    return reviews


# =========================================================
# ADMIN DASHBOARD
# =========================================================

@login_required(login_url='admin_login')
def admin_dashboard(request):

    # Only staff/admin
    if not request.user.is_staff:

        logout(request)

        return redirect(
            'admin_login'
        )


    # =====================================================
    # ALL REVIEWS
    # =====================================================

    all_reviews = Review.objects.all().order_by(
        '-created_at'
    )


    # =====================================================
    # TOTAL REVIEWS
    # =====================================================

    total_reviews = all_reviews.count()


    # =====================================================
    # UNIQUE STUDENTS
    # =====================================================

    total_students = (

        all_reviews

        .values(
            'student_email'
        )

        .distinct()

        .count()

    )


    # =====================================================
    # AVERAGE RATING
    # =====================================================

    average_rating = all_reviews.aggregate(
        avg_rating=Avg('rating')
    )['avg_rating']


    if average_rating is None:

        average_rating = 0


    # =====================================================
    # SENTIMENT COUNTS
    # =====================================================

    positive_count = all_reviews.filter(
        sentiment='positive'
    ).count()


    neutral_count = all_reviews.filter(
        sentiment='neutral'
    ).count()


    negative_count = all_reviews.filter(
        sentiment='negative'
    ).count()


    # =====================================================
    # RATING COUNTS
    # =====================================================

    rating_counts = []


    for rating in range(1, 6):

        count = all_reviews.filter(
            rating=rating
        ).count()

        rating_counts.append(
            count
        )


    # =====================================================
    # FILTERED REVIEWS
    # =====================================================

    filtered_reviews = get_filtered_reviews(
        request
    )


    # =====================================================
    # FILTER VALUES
    # =====================================================

    search_query = request.GET.get(
        'search',
        ''
    ).strip()


    selected_course = request.GET.get(
        'course',
        ''
    ).strip()


    selected_rating = request.GET.get(
        'rating',
        ''
    ).strip()


    selected_sentiment = request.GET.get(
        'sentiment',
        ''
    ).strip()


    # =====================================================
    # COURSE LIST
    # =====================================================

    courses = (

        Review.objects

        .values_list(
            'course',
            flat=True
        )

        .distinct()

        .order_by(
            'course'
        )

    )


    # =====================================================
    # PAGINATION
    # =====================================================

    paginator = Paginator(
        filtered_reviews,
        10
    )


    page_number = request.GET.get(
        'page'
    )


    reviews_page = paginator.get_page(
        page_number
    )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        # Overall statistics

        'total_reviews':
            total_reviews,

        'total_students':
            total_students,

        'average_rating':
            round(
                average_rating,
                2
            ),

        'positive_count':
            positive_count,

        'neutral_count':
            neutral_count,

        'negative_count':
            negative_count,

        'rating_counts':
            rating_counts,


        # Reviews

        'recent_reviews':
            reviews_page,


        # Filters

        'courses':
            courses,

        'search_query':
            search_query,

        'selected_course':
            selected_course,

        'selected_rating':
            selected_rating,

        'selected_sentiment':
            selected_sentiment,

    }


    return render(
        request,
        'admin_dashboard.html',
        context
    )


# =========================================================
# FULL REVIEW DETAIL
# =========================================================

@login_required(login_url='admin_login')
def review_detail(
    request,
    review_id
):

    # Only staff/admin
    if not request.user.is_staff:

        logout(request)

        return redirect(
            'admin_login'
        )


    try:

        review = Review.objects.get(
            id=review_id
        )


    except Review.DoesNotExist:

        return JsonResponse(
            {
                'error':
                    'Review not found.'
            },
            status=404
        )


    return JsonResponse({

        'id':
            review.id,

        'student_name':
            review.student_name,

        'student_email':
            review.student_email,

        'course':
            review.course,

        'batch':
            review.batch,

        'rating':
            review.rating,

        'teaching_quality':
            review.teaching_quality,

        'communication':
            review.communication,

        'practical_knowledge':
            review.practical_knowledge,

        'doubt_solving':
            review.doubt_solving,

        'sentiment':
            review.sentiment,

        'feedback':
            review.feedback,

        'is_anonymous':
            review.is_anonymous,

        'created_at':
            timezone.localtime(
                review.created_at
            ).strftime(
                '%d %b %Y, %I:%M %p'
            ),

    })


# =========================================================
# CSV EXPORT
# =========================================================

@login_required(login_url='admin_login')
def export_reviews_csv(request):

    # Only staff/admin
    if not request.user.is_staff:

        logout(request)

        return redirect(
            'admin_login'
        )


    # Get filtered reviews
    reviews = get_filtered_reviews(
        request
    )


    # Create response
    response = HttpResponse(
        content_type='text/csv'
    )


    response[
        'Content-Disposition'
    ] = (
        'attachment; '
        'filename="faculty_reviews.csv"'
    )


    writer = csv.writer(
        response
    )


    # =====================================================
    # HEADER
    # =====================================================

    writer.writerow([

        'Student Name',

        'Student Email',

        'Course',

        'Batch',

        'Overall Rating',

        'Teaching Quality',

        'Communication',

        'Practical Knowledge',

        'Doubt Solving',

        'Sentiment',

        'Feedback',

        'Submitted At',

    ])


    # =====================================================
    # DATA
    # =====================================================

    for review in reviews:

        writer.writerow([

            review.student_name,

            review.student_email,

            review.course,

            review.batch,

            review.rating,

            review.teaching_quality,

            review.communication,

            review.practical_knowledge,

            review.doubt_solving,

            review.sentiment,

            review.feedback,

            review.created_at.strftime(
                '%d %b %Y, %I:%M %p'
            ),

        ])


    return response


# =========================================================
# EXCEL EXPORT
# =========================================================

@login_required(login_url='admin_login')
def export_reviews_excel(request):

    # Only staff/admin
    if not request.user.is_staff:

        logout(request)

        return redirect(
            'admin_login'
        )


    # Get filtered reviews
    reviews = get_filtered_reviews(
        request
    )


    # =====================================================
    # CREATE WORKBOOK
    # =====================================================

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        "Faculty Reviews"
    )


    # =====================================================
    # HEADERS
    # =====================================================

    headers = [

        'Student Name',

        'Student Email',

        'Course',

        'Batch',

        'Overall Rating',

        'Teaching Quality',

        'Communication',

        'Practical Knowledge',

        'Doubt Solving',

        'Sentiment',

        'Feedback',

        'Submitted At',

    ]


    worksheet.append(
        headers
    )


    # =====================================================
    # HEADER STYLE
    # =====================================================

    header_fill = PatternFill(
        fill_type='solid',
        fgColor='2563EB'
    )


    header_font = Font(
        bold=True,
        color='FFFFFF'
    )


    header_alignment = Alignment(
        horizontal='center',
        vertical='center'
    )


    for cell in worksheet[1]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = (
            header_alignment
        )


    # =====================================================
    # DATA
    # =====================================================

    for review in reviews:

        worksheet.append([

            review.student_name,

            review.student_email,

            review.course,

            review.batch,

            review.rating,

            review.teaching_quality,

            review.communication,

            review.practical_knowledge,

            review.doubt_solving,

            review.sentiment,

            review.feedback,

            review.created_at.strftime(
                '%d %b %Y, %I:%M %p'
            ),

        ])


    # =====================================================
    # COLUMN WIDTHS
    # =====================================================

    column_widths = {

        'A': 22,

        'B': 32,

        'C': 20,

        'D': 15,

        'E': 16,

        'F': 20,

        'G': 18,

        'H': 22,

        'I': 18,

        'J': 15,

        'K': 45,

        'L': 25,

    }


    for column, width in (
        column_widths.items()
    ):

        worksheet.column_dimensions[
            column
        ].width = width


    # =====================================================
    # FREEZE HEADER
    # =====================================================

    worksheet.freeze_panes = (
        'A2'
    )


    # =====================================================
    # ALIGNMENT
    # =====================================================

    for row in worksheet.iter_rows(
        min_row=2
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical='top',
                wrap_text=True
            )


    # =====================================================
    # CREATE EXCEL FILE
    # =====================================================

    output = BytesIO()


    workbook.save(
        output
    )


    output.seek(0)


    # =====================================================
    # RESPONSE
    # =====================================================

    response = HttpResponse(

        output.getvalue(),

        content_type=(
            'application/'
            'vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )

    )


    response[
        'Content-Disposition'
    ] = (
        'attachment; '
        'filename="faculty_reviews.xlsx"'
    )


    return response


# =========================================================
# GOOGLE LOGIN
# =========================================================

def google_login(request):

    flow = create_google_flow()


    authorization_url, state = (
        flow.authorization_url(

            prompt='select_account',

            access_type='online'

        )
    )


    # Save OAuth state
    request.session[
        'google_oauth_state'
    ] = state


    # Save PKCE verifier
    request.session[
        'google_code_verifier'
    ] = flow.code_verifier


    return redirect(
        authorization_url
    )


# =========================================================
# GOOGLE CALLBACK
# =========================================================

def google_callback(request):

    # =====================================================
    # GET SAVED STATE
    # =====================================================

    saved_state = request.session.pop(
        'google_oauth_state',
        None
    )


    # =====================================================
    # GET PKCE VERIFIER
    # =====================================================

    code_verifier = request.session.pop(
        'google_code_verifier',
        None
    )


    if not saved_state:

        messages.error(
            request,
            'Google authentication session expired. Please try again.'
        )

        return redirect(
            'forgot_password'
        )


    if not code_verifier:

        messages.error(
            request,
            'Google security verification failed. Please try again.'
        )

        return redirect(
            'forgot_password'
        )


    try:

        # =================================================
        # CHECK STATE
        # =================================================

        returned_state = request.GET.get(
            'state'
        )


        if returned_state != saved_state:

            raise ValueError(
                'OAuth state mismatch.'
            )


        # =================================================
        # AUTHORIZATION CODE
        # =================================================

        authorization_code = (
            request.GET.get(
                'code'
            )
        )


        if not authorization_code:

            raise ValueError(
                'Authorization code missing.'
            )


        # =================================================
        # TOKEN EXCHANGE
        # =================================================

        token_response = requests.post(

            'https://oauth2.googleapis.com/token',

            data={

                'code':
                    authorization_code,

                'client_id':
                    settings.GOOGLE_CLIENT_ID,

                'client_secret':
                    settings.GOOGLE_CLIENT_SECRET,

                'redirect_uri':
                    settings.GOOGLE_REDIRECT_URI,

                'grant_type':
                    'authorization_code',

                'code_verifier':
                    code_verifier,

            },

            timeout=20

        )


        print(
            'Google token status:',
            token_response.status_code
        )


        if not token_response.ok:

            print(
                'Google token response:',
                token_response.text
            )


            raise ValueError(
                'Google token exchange failed.'
            )


        # =================================================
        # TOKEN DATA
        # =================================================

        token_data = (
            token_response.json()
        )


        google_id_token = (
            token_data.get(
                'id_token'
            )
        )


        if not google_id_token:

            raise ValueError(
                'Google ID token was not returned.'
            )


        # =================================================
        # VERIFY GOOGLE ID TOKEN
        # =================================================

        google_user = (
            id_token.verify_oauth2_token(

                google_id_token,

                google_requests.Request(),

                settings.GOOGLE_CLIENT_ID,

                clock_skew_in_seconds=60

            )
        )


    except Exception as e:

        print(
            '================================'
        )

        print(
            'GOOGLE AUTH ERROR:',
            repr(e)
        )

        print(
            '================================'
        )


        messages.error(
            request,
            'Google authentication failed.'
        )


        return redirect(
            'forgot_password'
        )


    # =====================================================
    # GOOGLE EMAIL
    # =====================================================

    google_email = (
        google_user.get(
            'email',
            ''
        )
        .strip()
        .lower()
    )


    # =====================================================
    # EMAIL VERIFIED?
    # =====================================================

    email_verified = (
        google_user.get(
            'email_verified',
            False
        )
    )


    # =====================================================
    # ADMIN EMAIL
    # =====================================================

    admin_email = (
        settings.ADMIN_EMAIL
        .strip()
        .lower()
    )


    if not email_verified:

        messages.error(
            request,
            'Google email is not verified.'
        )

        return redirect(
            'forgot_password'
        )


    # =====================================================
    # CHECK ADMIN EMAIL
    # =====================================================

    if google_email != admin_email:

        messages.error(
            request,
            'This Google account is not authorized for admin recovery.'
        )

        return redirect(
            'forgot_password'
        )


    # =====================================================
    # GOOGLE VERIFIED
    # =====================================================

    request.session[
        'google_verified_for_reset'
    ] = True


    return redirect(
        'reset_password'
    )


# =========================================================
# FORGOT PASSWORD
# =========================================================

def forgot_password(request):

    return render(
        request,
        'forgot_password.html'
    )


# =========================================================
# RESET PASSWORD
# =========================================================

def reset_password(request):

    # =====================================================
    # GOOGLE VERIFICATION REQUIRED
    # =====================================================

    if not request.session.get(
        'google_verified_for_reset'
    ):

        messages.error(
            request,
            'Please verify your Google account first.'
        )

        return redirect(
            'forgot_password'
        )


    # =====================================================
    # POST
    # =====================================================

    if request.method == 'POST':

        password = request.POST.get(
            'password'
        )


        confirm_password = request.POST.get(
            'confirm_password'
        )


        # =================================================
        # PASSWORD MATCH
        # =================================================

        if password != confirm_password:

            messages.error(
                request,
                'Passwords do not match.'
            )


            return render(
                request,
                'reset_password.html'
            )


        # =================================================
        # PASSWORD VALIDATION
        # =================================================

        try:

            validate_password(
                password
            )


        except ValidationError as e:

            messages.error(
                request,
                ' '.join(
                    e.messages
                )
            )


            return render(
                request,
                'reset_password.html'
            )


        # =================================================
        # FIND ADMIN
        # =================================================

        try:

            admin_user = User.objects.get(

                email__iexact=
                    settings.ADMIN_EMAIL,

                is_staff=True

            )


            # Set new password
            admin_user.set_password(
                password
            )


            admin_user.save()


            # Remove Google verification
            request.session.pop(
                'google_verified_for_reset',
                None
            )


            messages.success(
                request,
                'Password reset successfully. Please login with your new password.'
            )


            return redirect(
                'admin_login'
            )


        except User.DoesNotExist:

            messages.error(
                request,
                'Admin account not found.'
            )


            return redirect(
                'forgot_password'
            )


    return render(
        request,
        'reset_password.html'
    )
# =========================================================
# MANAGE PROFILE
# =========================================================

@login_required(login_url='admin_login')
def manage_profile(request):

    # =====================================================
    # ADMIN ACCESS ONLY
    # =====================================================

    if not request.user.is_staff:

        logout(request)

        return redirect('admin_login')


    # =====================================================
    # GET EXISTING PROFILE
    # =====================================================

    profile = FacultyProfile.objects.first()


    # =====================================================
    # POST
    # =====================================================

    if request.method == 'POST':

        cropped_image_data = request.POST.get(
            'cropped_image',
            ''
        ).strip()


        # =================================================
        # MAIN PROFILE FORM
        # =================================================

        if cropped_image_data:

            # IMPORTANT:
            # Do not pass request.FILES here.
            # We want to save the CROPPED image,
            # not the original uploaded image.

            form = FacultyProfileForm(
                request.POST,
                instance=profile
            )

        else:

            form = FacultyProfileForm(
                request.POST,
                request.FILES,
                instance=profile
            )


        # =================================================
        # FORM VALIDATION
        # =================================================

        if form.is_valid():

            try:

                with transaction.atomic():

                    profile = form.save(
                        commit=False
                    )


                    # =====================================
                    # SAVE CROPPED IMAGE
                    # =====================================

                    if cropped_image_data:

                        try:

                            # ---------------------------------
                            # Remove data:image/jpeg;base64,
                            # ---------------------------------

                            if ',' in cropped_image_data:

                                header, encoded = (
                                    cropped_image_data.split(
                                        ',',
                                        1
                                    )
                                )

                            else:

                                encoded = cropped_image_data


                            # ---------------------------------
                            # Decode base64
                            # ---------------------------------

                            image_bytes = base64.b64decode(
                                encoded,
                                validate=True
                            )


                            # ---------------------------------
                            # Basic safety/size check
                            # ---------------------------------

                            if not image_bytes:

                                raise ValueError(
                                    'Empty cropped image.'
                                )


                            # ---------------------------------
                            # Keep old image path
                            # ---------------------------------

                            old_image_name = None

                            if (
                                profile.pk
                                and profile.profile_image
                            ):

                                old_image_name = (
                                    profile.profile_image.name
                                )


                            # ---------------------------------
                            # Create new image file
                            # ---------------------------------

                            new_image = ContentFile(
                                image_bytes,
                                name=(
                                    f"faculty_profile_"
                                    f"{request.user.id}.jpg"
                                )
                            )


                            # ---------------------------------
                            # Assign image
                            # ---------------------------------

                            profile.profile_image = new_image


                            # ---------------------------------
                            # Save profile
                            # ---------------------------------

                            profile.save()


                            # ---------------------------------
                            # Delete old image
                            # ---------------------------------

                            if (
                                old_image_name
                                and default_storage.exists(
                                    old_image_name
                                )
                                and old_image_name != (
                                    profile.profile_image.name
                                )
                            ):

                                default_storage.delete(
                                    old_image_name
                                )


                        except (
                            ValueError,
                            binascii.Error
                        ) as e:

                            messages.error(
                                request,
                                'Invalid cropped image. '
                                'Please crop the image again.'
                            )

                            raise e


                    else:

                        # =================================
                        # NO CROPPED IMAGE
                        # =================================

                        profile.save()


                    # =====================================
                    # EDUCATION FORMSET
                    # =====================================

                    education_formset = EducationFormSet(
                        request.POST,
                        instance=profile,
                        prefix='education'
                    )


                    # =====================================
                    # SKILL FORMSET
                    # =====================================

                    skill_formset = SkillFormSet(
                        request.POST,
                        instance=profile,
                        prefix='skill'
                    )


                    # =====================================
                    # VALIDATE FORMSETS
                    # =====================================

                    if (
                        education_formset.is_valid()
                        and
                        skill_formset.is_valid()
                    ):

                        education_formset.save()

                        skill_formset.save()

                    else:

                        # Re-render with errors
                        return render(
                            request,
                            'manage_profile.html',
                            {
                                'form': form,
                                'profile': profile,
                                'education_formset':
                                    education_formset,
                                'skill_formset':
                                    skill_formset,
                            }
                        )


                # =========================================
                # SUCCESS
                # =========================================

                messages.success(
                    request,
                    'Profile updated successfully.'
                )

                return redirect(
                    'manage_profile'
                )


            except Exception as e:

                print(
                    'PROFILE SAVE ERROR:',
                    e
                )

                messages.error(
                    request,
                    'Something went wrong while saving '
                    'the profile.'
                )


    # =====================================================
    # GET
    # =====================================================

    else:

        if profile:

            form = FacultyProfileForm(
                instance=profile
            )

            education_formset = EducationFormSet(
                instance=profile,
                prefix='education'
            )

            skill_formset = SkillFormSet(
                instance=profile,
                prefix='skill'
            )

        else:

            form = FacultyProfileForm()

            blank_profile = FacultyProfile()

            education_formset = EducationFormSet(
                instance=blank_profile,
                prefix='education'
            )

            skill_formset = SkillFormSet(
                instance=blank_profile,
                prefix='skill'
            )


    # =====================================================
    # RENDER
    # =====================================================

    return render(
        request,
        'manage_profile.html',
        {
            'form': form,
            'profile': profile,
            'education_formset':
                education_formset,
            'skill_formset':
                skill_formset,
        }
    )